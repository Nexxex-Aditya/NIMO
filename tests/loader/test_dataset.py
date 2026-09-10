"""Loader tests against the real workbook, plus fixture-based drift tests.

`specs/loader.md` "Tests": whole-dataset checks run against the real committed
`data/raw/` file, because a synthetic fixture cannot tell you whether the real
defect is still being caught. Assertions that must *fire* are exercised
against mutated copies in `tmp_path` — `data/raw/` is never written (`04` §12).
"""

import shutil
import socket
from pathlib import Path

import openpyxl
import pytest

from nimo.contracts import CharacteristicRule, RawRow
from nimo.loader import (
    applicable_characteristics,
    assert_guideline_modules_subset,
    characteristic_rule,
    load_characteristic_guidelines,
    load_characteristic_rules,
    load_qa_header,
    load_rows,
)
from nimo.loader.dataset import (
    EXPECTED_DEV_CORRUPT_BARCODES,
    EXPECTED_DEV_DISTINCT_BARCODES,
    EXPECTED_DEV_ROWS,
    EXPECTED_HEADER,
)
from nimo.loader.errors import DatasetDriftError, DatasetSchemaError, RetailerNotMappedError

REPO_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = REPO_ROOT / "data" / "raw" / "product_truth_agent_dataset.xlsx"
RETAILERS = REPO_ROOT / "config" / "retailers.yaml"


@pytest.fixture(scope="module")
def dev_rows() -> list[RawRow]:
    return load_rows(WORKBOOK, "dev", RETAILERS)


@pytest.fixture(scope="module")
def qa_rows() -> list[RawRow]:
    return load_rows(WORKBOOK, "qa", RETAILERS)


@pytest.fixture(scope="module")
def rules() -> list[CharacteristicRule]:
    return load_characteristic_rules(WORKBOOK)


def _copy_workbook(tmp_path: Path) -> Path:
    """A writable copy. `data/raw/` is read-only, always (`04` §12)."""
    destination = tmp_path / "mutated.xlsx"
    shutil.copy(WORKBOOK, destination)
    return destination


def _mutate_cell(path: Path, sheet: str, row: int, column: int, value: str) -> None:
    workbook = openpyxl.load_workbook(path)
    workbook[sheet].cell(row=row, column=column, value=value)
    workbook.save(path)
    workbook.close()


# --- Regression, exact: the P2 gate (`specs/loader.md` §11) ---------------


def test_dev_barcode_corruption_count_is_exact(dev_rows: list[RawRow]) -> None:
    """`01` §3: 377 of 412. Asserted against the real file, not a fixture."""
    assert len(dev_rows) == EXPECTED_DEV_ROWS
    assert sum(row.barcode_corrupt for row in dev_rows) == EXPECTED_DEV_CORRUPT_BARCODES


def test_dev_distinct_barcode_count_is_exact(dev_rows: list[RawRow]) -> None:
    distinct = {row.barcode_raw for row in dev_rows if row.barcode_raw is not None}
    assert len(distinct) == EXPECTED_DEV_DISTINCT_BARCODES


def test_every_corrupt_dev_barcode_is_nulled(dev_rows: list[RawRow]) -> None:
    """The registry-safety invariant, checked across the whole real sheet."""
    for row in dev_rows:
        if row.barcode_corrupt:
            assert row.barcode is None
            assert row.barcode_raw is not None


def test_qa_barcodes_are_clean(qa_rows: list[RawRow]) -> None:
    """`01` §3: qa is intact — the asymmetry that makes dev-only tuning a trap."""
    assert len(qa_rows) == EXPECTED_DEV_ROWS
    assert sum(row.barcode_corrupt for row in qa_rows) == 0
    assert all(row.barcode is not None for row in qa_rows)


def test_drift_detection_fires_when_corruption_count_changes(tmp_path: Path) -> None:
    """Proves criterion 2 actually fires, not that the number sits unused.

    One corrupt dev row is "repaired" in a copy, taking the count to 376.
    """
    mutated = _copy_workbook(tmp_path)
    external_code_column = EXPECTED_HEADER.index("EXTERNAL_CODE") + 1
    _mutate_cell(mutated, "dev", 2, external_code_column, "5014697056627")
    with pytest.raises(DatasetDriftError, match="corruption count"):
        load_rows(mutated, "dev", RETAILERS)


# --- §4 retailer lookup --------------------------------------------------


def test_boots_homescan_sentinel_resolves_to_boots(
    dev_rows: list[RawRow], qa_rows: list[RawRow]
) -> None:
    """The mandatory named regression (`01` §12, `05` §5).

    A naive regex captures "(HOMESCAN)" here and silently discards BOOTS —
    the largest retailer in the dataset. It must never regress.
    """
    sentinel = [row for row in dev_rows + qa_rows if row.retailer_raw == "BOOTS (GB) (HOMESCAN)"]
    assert sentinel, "the sentinel row disappeared from the dataset"
    for row in sentinel:
        assert row.retailer == "BOOTS"
        assert row.retailer != "(HOMESCAN)"


def test_retailer_table_covers_every_value_in_the_real_data(
    dev_rows: list[RawRow], qa_rows: list[RawRow]
) -> None:
    """All 50 entries present — the P2 precondition, checked against reality."""
    distinct = {row.retailer_raw for row in dev_rows + qa_rows}
    assert len(distinct) == 50
    assert all(row.retailer for row in dev_rows + qa_rows)


def test_retailer_raw_is_preserved_verbatim(dev_rows: list[RawRow]) -> None:
    """§4: `retailer_raw` stays the canonical key for both lookups."""
    row = next(r for r in dev_rows if r.retailer_raw == "P00R4 (GB) BOOTS")
    assert row.retailer_raw == "P00R4 (GB) BOOTS"
    assert row.retailer == "BOOTS"


def test_unmapped_retailer_raises_rather_than_falling_back(tmp_path: Path) -> None:
    """§4: fail loud. A silent fallback to the raw string is the failure class
    that already cost 49 rows' worth of correct names once."""
    partial = tmp_path / "partial_retailers.yaml"
    partial.write_text('"OCADO":\n  name: OCADO\n  domain: ocado.com\n', encoding="utf-8")
    with pytest.raises(RetailerNotMappedError, match="has no entry"):
        load_rows(WORKBOOK, "dev", partial)


def test_retailer_entry_without_a_name_raises(tmp_path: Path) -> None:
    """`name` is complete-or-nothing; `domain` may stay partial until P7."""
    nameless = tmp_path / "nameless.yaml"
    nameless.write_text('"OCADO":\n  domain: ocado.com\n', encoding="utf-8")
    with pytest.raises(DatasetSchemaError, match="has no `name`"):
        load_rows(WORKBOOK, "dev", nameless)


# --- §2a encoding repair, on the real file -------------------------------


def test_real_mojibake_rows_are_repaired_and_flagged(
    dev_rows: list[RawRow], qa_rows: list[RawRow]
) -> None:
    """`01` §13: the JASÖN rows. Repair happens *and* is never silent."""
    repaired = [row for row in dev_rows + qa_rows if row.brand_encoding_suspect]
    assert len(repaired) == 3
    for row in repaired:
        assert row.brand_raw == "JASÖN"
        assert row.brand == "JASÖN"
        assert "Ã" not in row.brand_raw


def test_legitimate_multilingual_rows_are_not_flagged(
    dev_rows: list[RawRow], qa_rows: list[RawRow]
) -> None:
    """Unconditional ftfy must not "repair" text that was already correct."""
    for row in dev_rows + qa_rows:
        if "pärla" in row.desc_raw or "antibactérien" in row.desc_raw:
            assert row.desc_encoding_suspect is False


# --- §8 char_value_list --------------------------------------------------


def test_characteristic_rules_shape(rules: list[CharacteristicRule]) -> None:
    assert len(rules) == 195
    assert len({rule.module for rule in rules}) == 59
    assert len({rule.characteristic for rule in rules}) == 13


def test_possible_values_parse_to_atomic_lists(rules: list[CharacteristicRule]) -> None:
    """§8: `ast.literal_eval`, never `eval`, never string-splitting."""
    bristle = characteristic_rule(
        rules, "ORAL HYGIENE - COMBINATION PACKS", "GLOBAL_BRISTLE_STRENGTH_CLAIM"
    )
    assert bristle is not None
    assert bristle.allowed_values == ["HARD", "MEDIUM", "NO CLAIM", "SOFT"]
    assert bristle.open_close == "Close"
    assert bristle.binary is False

    percentage = characteristic_rule(
        rules, "TONGUE CLEANING - BRUSH/SCRAPER", "GLOBAL_PERCENTAGE_NATURAL_INGREDIENTS"
    )
    assert percentage is not None
    assert percentage.allowed_values == ["100%", "95%"]
    assert percentage.open_close == "Open-ended"


def test_allowed_values_for_combinator_characteristics_are_atomic(
    rules: list[CharacteristicRule],
) -> None:
    """`01` §11: `&` is a combinator in the *values*, never inside a vocabulary
    entry. The loader exposes the atoms; P12 does the component-wise check."""
    for name in ("GLOBAL_ORAL_CARE_FUNCTION", "GLOBAL_CONSUMER_LIFESTAGE_CLAIM"):
        matching = [rule for rule in rules if rule.characteristic == name]
        assert matching, f"{name} absent from char_value_list"
        for rule in matching:
            for value in rule.allowed_values:
                assert "&" not in value, f"{name} vocabulary entry {value!r} contains '&'"


def test_non_oral_health_category_raises(tmp_path: Path) -> None:
    """Single-category scope is a standing assumption across `03`/`04`/`05`."""
    mutated = _copy_workbook(tmp_path)
    _mutate_cell(mutated, "char_value_list", 2, 1, "PERSONAL CARE")
    with pytest.raises(DatasetSchemaError, match="category"):
        load_characteristic_rules(mutated)


# --- §7/§9 characteristic-name mapping -----------------------------------


def test_both_rule_sheets_map_bijectively_onto_the_dev_qa_columns(
    rules: list[CharacteristicRule],
) -> None:
    """§7: total and bijective, after the alias table is applied."""
    expected = {name for name in EXPECTED_HEADER if name.startswith("GLOBAL_")}
    assert len(expected) == 13
    assert {rule.characteristic for rule in rules} == expected
    guidelines = load_characteristic_guidelines(WORKBOOK)
    assert {guideline.characteristic for guideline in guidelines} == expected


def test_renamed_characteristic_raises(tmp_path: Path) -> None:
    """Proves the bijectivity assertion fires, rather than merely not tripping
    on clean input."""
    mutated = _copy_workbook(tmp_path)
    # char_guidelines column 3 is CHARACTERISTICS NAME
    _mutate_cell(mutated, "char_guidelines", 2, 3, "GLOBAL INVENTED CLAIM")
    with pytest.raises(DatasetSchemaError, match="bijectiv"):
        load_characteristic_guidelines(mutated)


def test_guideline_modules_are_a_subset_of_rule_modules(
    rules: list[CharacteristicRule],
) -> None:
    """§9 cross-sheet consistency check."""
    guidelines = load_characteristic_guidelines(WORKBOOK)
    assert len(guidelines) == 196
    assert_guideline_modules_subset(guidelines, rules)


def test_interspace_is_applicable_to_the_interdental_module(
    rules: list[CharacteristicRule],
) -> None:
    """`01` §6: null in all 412 dev rows, but NOT a dead column.

    It applies to exactly one module, which dev happens not to cover — and two
    real qa rows need it predicted. Reading the dev null-rate as "never
    applicable" was a documented past mistake; this pins the correction.
    """
    applicable = applicable_characteristics(rules, "TOOTHBRUSHES - MANUAL - INTERDENTAL")
    assert "GLOBAL_INTERSPACE_CLAIM" in applicable


def test_applicable_characteristics_excludes_non_applicable(
    rules: list[CharacteristicRule],
) -> None:
    """Filling a non-applicable characteristic is scored wrong, not partial."""
    toothpaste = applicable_characteristics(
        rules, "TOOTH CLEANING - FOAM/GEL/LIQUID/PASTE (NATURAL TEETH)"
    )
    assert "GLOBAL_BRISTLE_STRENGTH_CLAIM" not in toothpaste
    assert "GLOBAL_IF_WITH_FLUORIDE" in toothpaste


# --- §10 qa header -------------------------------------------------------


def test_qa_header_matches_the_submission_contract() -> None:
    header = load_qa_header(WORKBOOK)
    assert header == EXPECTED_HEADER
    assert len(header) == 23


def test_reordered_qa_header_raises(tmp_path: Path) -> None:
    """A silently reordered submission scores zero regardless of values."""
    mutated = _copy_workbook(tmp_path)
    _mutate_cell(mutated, "qa", 1, 1, "NAN_KEY")
    _mutate_cell(mutated, "qa", 1, 2, "ITEM_CODE")
    with pytest.raises(DatasetDriftError, match="qa header"):
        load_qa_header(mutated)


def test_renamed_qa_column_raises(tmp_path: Path) -> None:
    mutated = _copy_workbook(tmp_path)
    _mutate_cell(mutated, "qa", 1, 1, "ITEM_ID")
    with pytest.raises(DatasetDriftError, match="qa header"):
        load_qa_header(mutated)


def test_missing_sheet_raises(tmp_path: Path) -> None:
    mutated = _copy_workbook(tmp_path)
    workbook = openpyxl.load_workbook(mutated)
    del workbook["qa"]
    workbook.save(mutated)
    workbook.close()
    with pytest.raises(DatasetDriftError, match="absent"):
        load_qa_header(mutated)


# --- Determinism (`04` §5) ----------------------------------------------


def test_row_order_matches_source_order(dev_rows: list[RawRow]) -> None:
    """§Determinism: no sorting, no groupby reordering."""
    import pandas as pd

    frame = pd.read_excel(WORKBOOK, sheet_name="dev", dtype=str)
    assert [row.nan_key for row in dev_rows] == [int(str(v)) for v in frame["NAN_KEY"]]


def test_loading_twice_is_byte_identical(dev_rows: list[RawRow]) -> None:
    again = load_rows(WORKBOOK, "dev", RETAILERS)
    assert [row.model_dump_json() for row in again] == [row.model_dump_json() for row in dev_rows]


# --- `04` §6/§8: no network ---------------------------------------------


def test_no_network_access_during_load(monkeypatch: pytest.MonkeyPatch) -> None:
    """Asserted rather than assumed, per the standing pattern (`04` §8)."""

    def _forbidden(*args: object, **kwargs: object) -> None:
        raise AssertionError("the loader must not open a socket")

    monkeypatch.setattr(socket, "socket", _forbidden)
    monkeypatch.setattr(socket, "create_connection", _forbidden)
    rows = load_rows(WORKBOOK, "qa", RETAILERS)
    assert len(rows) == 412


# --- Row identity (`01` §14) --------------------------------------------


def test_row_uid_is_unique_where_nan_key_is_not(
    dev_rows: list[RawRow], qa_rows: list[RawRow]
) -> None:
    """`01` §14: NAN_KEY collides across genuinely different products.

    15 dev NAN_KEYs are duplicated (42 rows) and 16 in qa (46 rows), every
    one of them a rounding artifact — and 11 dev NAN_KEYs span more than one
    MODULE, i.e. unrelated products sharing a key. `03` §2/§5 key per-row
    artifacts and resumability on the row identity, so it has to be unique or
    one product's cached result is served for another, silently.
    """
    for name, rows in (("dev", dev_rows), ("qa", qa_rows)):
        uids = [row.row_uid for row in rows]
        assert len(set(uids)) == len(rows), f"{name} row_uid is not unique"
        nan_keys = [row.nan_key for row in rows]
        assert len(set(nan_keys)) < len(rows), (
            f"{name} NAN_KEY is unexpectedly unique — `01` §14 records it as colliding. "
            f"If the organizers fixed it, re-verify 01 before relying on it."
        )


def test_row_uid_encodes_sheet_and_position(dev_rows: list[RawRow]) -> None:
    """Positional and deterministic — same file in, same identity out."""
    assert dev_rows[0].row_uid == "dev:0"
    assert dev_rows[411].row_uid == "dev:411"


def test_row_uid_does_not_collide_across_sheets(
    dev_rows: list[RawRow], qa_rows: list[RawRow]
) -> None:
    assert not ({row.row_uid for row in dev_rows} & {row.row_uid for row in qa_rows})
