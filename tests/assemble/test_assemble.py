"""P14 tests — `specs/assemble.md` §6. `01` §10 #7, `04` §4/§5, `05` §5.

The artifact tree comes from the real offline runner over a few dev rows —
the same code path production uses — never from hand-built fixtures
(decision log 2026-09-11: tests that build fixtures by hand verify the
function, not the wiring).
"""

import json
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from nimo.assemble import (
    AssemblyError,
    AssemblyReport,
    OutputConfig,
    OutputConfigError,
    assemble_rows,
    format_report,
    load_output_config,
    to_csv_text,
    write_csv,
    write_xlsx,
)
from nimo.contracts import CharacteristicRule, ModulePrediction, OutputRow, ProductQuery, RawRow
from nimo.loader import load_characteristic_rules, read_header
from nimo.run import RunPaths, Stages, artifact_path, run
from nimo.run.artifacts import STAGE_SEQUENCE
from tests.run.test_runner import WORKBOOK, fixed_clock

CONFIG = load_output_config()
N = 6  # dev rows driven through the real offline runner


@pytest.fixture(scope="module")
def rules() -> list[CharacteristicRule]:
    return load_characteristic_rules(WORKBOOK)


@pytest.fixture(scope="module")
def tree(tmp_path_factory: pytest.TempPathFactory, dev_rows: list[RawRow], stages: Stages) -> Path:
    """Six dev rows through the runner; row dev:3 deliberately failed at
    classify so the tree has one incomplete row."""
    root = tmp_path_factory.mktemp("assemble")
    paths = RunPaths(
        artifacts=root / "artifacts",
        trace=root / "trace.jsonl",
        failures=root / "failures.jsonl",
        config_dir=WORKBOOK.parent.parent.parent / "config",
    )

    def failing_classify(query: ProductQuery) -> ModulePrediction:
        if query.row_uid == "dev:3":
            raise RuntimeError("deliberate")
        return stages.classify(query)

    run(dev_rows[:N], replace(stages, classify=failing_classify), paths, "r", fixed_clock)
    return root


def assembled(
    tree: Path, rules: list[CharacteristicRule], config: OutputConfig = CONFIG
) -> tuple[list[OutputRow], AssemblyReport]:
    return assemble_rows(
        WORKBOOK, "dev", tree / "artifacts", rules, config, tree / "failures.jsonl"
    )


# --- header and shape ----------------------------------------------------------------


def test_the_header_is_the_workbooks_exactly() -> None:
    assert (
        list(OutputRow.model_fields) == read_header(WORKBOOK, "qa") == read_header(WORKBOOK, "dev")
    )


def test_every_sheet_row_is_emitted_in_order(tree: Path, rules: list[CharacteristicRule]) -> None:
    rows, report = assembled(tree, rules)
    assert len(rows) == 412 and report.rows == 412
    assert report.complete == N - 1
    header = read_header(WORKBOOK, "dev")
    assert to_csv_text(rows).splitlines()[0] == ",".join(header)


# --- complete rows, blank rows, passthrough ----------------------------------------------


def test_a_complete_row_carries_the_artifact_values(
    tree: Path, rules: list[CharacteristicRule]
) -> None:
    rows, _ = assembled(tree, rules)
    row = rows[0]
    module = json.loads(
        artifact_path(tree / "artifacts", "classify", "dev:0").read_text(encoding="utf-8")
    )
    reasoning = json.loads(
        artifact_path(tree / "artifacts", "reason", "dev:0").read_text(encoding="utf-8")
    )
    assert module["module"] == row.MODULE
    assert reasoning["text"] == row.REASONING
    assert row.PRODUCT_URL is None  # offline: abstained


def test_a_failed_row_is_blank_in_every_output_column(
    tree: Path, rules: list[CharacteristicRule]
) -> None:
    """`04` §4: never a partial row. Inputs pass through, outputs are empty."""
    rows, report = assembled(tree, rules)
    failed = rows[3]
    assert failed.MODULE is None and failed.REASONING is None and failed.PRODUCT_URL is None
    assert all(
        getattr(failed, name) is None
        for name in OutputRow.model_fields
        if name.startswith("GLOBAL_")
    )
    assert failed.RETAILER_DESC and failed.BRAND and failed.EXTERNAL_CODE
    assert report.blank_row_uids[:1] == ["dev:3"]
    assert report.failure_stages["dev:3"] == "classify"


def test_passthrough_is_the_workbooks_bytes_not_the_repaired_row(
    tree: Path, rules: list[CharacteristicRule]
) -> None:
    """`01` §13: three dev rows carry the mojibake brand `JASÃƒâ€“N`. The
    loader repairs it for the pipeline; the submission carries the
    organizers' own text."""
    rows, _ = assembled(tree, rules)
    import pandas as pd

    frame = pd.read_excel(WORKBOOK, sheet_name="dev", dtype=str)
    for index in range(412):
        assert frame.at[index, "BRAND"] == rows[index].BRAND
        assert frame.at[index, "RETAILER_DESC"] == rows[index].RETAILER_DESC
    mojibake = [r for r in rows if "Ã" in r.BRAND]
    assert len(mojibake) == 3


def test_external_code_and_keys_are_text_in_the_xlsx(
    tree: Path, rules: list[CharacteristicRule], tmp_path: Path
) -> None:
    """`01` §3 / `05` §5: the defect this project started with must not be
    reintroduced by our own write."""
    rows, _ = assembled(tree, rules)
    path = tmp_path / "s.xlsx"
    write_xlsx(rows, path, "dev", CONFIG)
    sheet = openpyxl.load_workbook(path)["dev"]
    header = [cell.value for cell in sheet[1]]
    assert header == list(OutputRow.model_fields)
    for name in ("ITEM_CODE", "NAN_KEY", "EXTERNAL_CODE"):
        cell = sheet.cell(row=2, column=header.index(name) + 1)
        assert cell.data_type == "s" and cell.number_format == "@", name
    assert sheet.cell(row=2, column=3).value == rows[0].EXTERNAL_CODE


# --- validation (`03` §4 stage 8) -------------------------------------------------------


def test_a_module_outside_the_set_raises(tree: Path, rules: list[CharacteristicRule]) -> None:
    path = artifact_path(tree / "artifacts", "classify", "dev:1")
    original = path.read_text(encoding="utf-8")
    try:
        path.write_text(
            original.replace(json.loads(original)["module"], "NOT A MODULE"), encoding="utf-8"
        )
        with pytest.raises(AssemblyError, match="not in char_value_list"):
            assembled(tree, rules)
    finally:
        path.write_text(original, encoding="utf-8")


def test_a_non_applicable_characteristic_value_raises(
    tree: Path, rules: list[CharacteristicRule]
) -> None:
    path = artifact_path(tree / "artifacts", "characteristics", "dev:1")
    original = path.read_text(encoding="utf-8")
    data = json.loads(original)
    module = data["module"]
    not_applicable = next(name for name in data["values"] if name not in data["applicable"])
    data["values"][not_applicable] = "PLASTIC"
    try:
        path.write_text(json.dumps(data), encoding="utf-8")
        with pytest.raises(AssemblyError, match="not applicable"):
            assembled(tree, rules)
    finally:
        path.write_text(original, encoding="utf-8")
    assert module  # the module was real; only the value was out of place


def test_a_misaligned_artifact_raises(tree: Path, rules: list[CharacteristicRule]) -> None:
    """`01` §14 at the output boundary: an artifact whose keys differ from
    the sheet row's must never attach its answer to that row."""
    path = artifact_path(tree / "artifacts", "normalize", "dev:1")
    original = path.read_text(encoding="utf-8")
    data = json.loads(original)
    data["nan_key"] = data["nan_key"] + 1
    try:
        path.write_text(json.dumps(data), encoding="utf-8")
        with pytest.raises(AssemblyError, match="misaligned"):
            assembled(tree, rules)
    finally:
        path.write_text(original, encoding="utf-8")


def test_product_url_field_switch_emits_the_title(
    tree: Path, rules: list[CharacteristicRule]
) -> None:
    """`[PROVISIONAL — Q2]`: a config value, not a module."""
    path = artifact_path(tree / "artifacts", "match", "dev:1")
    original = path.read_text(encoding="utf-8")
    data = json.loads(original)
    data["url"], data["page_title"] = "https://boots.com/p", "Aquafresh | Boots"
    try:
        path.write_text(json.dumps(data), encoding="utf-8")
        by_url, _ = assembled(tree, rules)
        by_title, _ = assembled(tree, rules, replace(CONFIG, product_url_field="title"))
        assert by_url[1].PRODUCT_URL == "https://boots.com/p"
        assert by_title[1].PRODUCT_URL == "Aquafresh | Boots"
    finally:
        path.write_text(original, encoding="utf-8")


# --- byte identity (`04` §5) -------------------------------------------------------------


def test_twice_run_files_are_byte_identical(
    tree: Path, rules: list[CharacteristicRule], tmp_path: Path
) -> None:
    rows, _ = assembled(tree, rules)
    first_csv, first_xlsx = tmp_path / "a.csv", tmp_path / "a.xlsx"
    second_csv, second_xlsx = tmp_path / "b.csv", tmp_path / "b.xlsx"
    write_csv(rows, first_csv)
    write_xlsx(rows, first_xlsx, "dev", CONFIG)
    write_csv(rows, second_csv)
    write_xlsx(rows, second_xlsx, "dev", CONFIG)
    assert first_csv.read_bytes() == second_csv.read_bytes()
    assert first_xlsx.read_bytes() == second_xlsx.read_bytes()


def test_report_names_blank_rows_and_fill_counts(
    tree: Path, rules: list[CharacteristicRule]
) -> None:
    _, report = assembled(tree, rules)
    text = format_report(report)
    assert "412 rows, 5 complete, 407 blank" in text
    assert "dev:3(classify)" in text and "PRODUCT_URL carries: url" in text
    assert "MODULE" in text


def test_config_rejects_an_unknown_product_url_field(tmp_path: Path) -> None:
    path = tmp_path / "output.yaml"
    path.write_text(
        "product_url_field: both\nxlsx_sheet_name_from_source: true\n"
        "fixed_timestamp: '2026-09-11T00:00:00Z'\n",
        encoding="utf-8",
    )
    with pytest.raises(OutputConfigError, match="product_url_field"):
        load_output_config(path)


def test_stage_sequence_is_what_assembly_reads() -> None:
    assert STAGE_SEQUENCE[-4:] == ("match", "classify", "characteristics", "reason")
