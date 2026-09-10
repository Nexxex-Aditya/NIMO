"""Workbook reading and row assembly — `specs/loader.md` §1, §8–§11.

Reads `data/raw/product_truth_agent_dataset.xlsx` (read-only, always — `04`
§12) and produces the four things `specs/loader.md` "Scope" names:
`list[RawRow]` per sheet, `list[CharacteristicRule]`,
`list[CharacteristicGuideline]`, and the captured `qa_header`.
"""

import ast
from pathlib import Path
from typing import Literal

import openpyxl
import pandas as pd
import structlog
import yaml
from openpyxl.worksheet.worksheet import Worksheet

from nimo.contracts import CharacteristicGuideline, CharacteristicRule, OutputRow, RawRow
from nimo.loader.errors import DatasetDriftError, DatasetSchemaError, RetailerNotMappedError
from nimo.loader.fields import (
    collapse_whitespace,
    normalize_characteristic_name,
    parse_barcode,
    parse_brand,
    parse_countries,
    repair_encoding,
)

log = structlog.get_logger(__name__)

# --- Dataset fingerprints, `01` §1/§3 -----------------------------------
# Deliberately module constants rather than `config/` values. `04` §9's
# zero-magic-numbers rule targets *tunables* — weights, thresholds, rate
# limits — things you are meant to change. These are the opposite: assertions
# that the source file still is what `01` documents. `specs/loader.md` §11 is
# explicit that the right response to a mismatch is to re-verify `01`, "don't
# just update this number", and moving them into config would make casually
# editing them easier, not harder.
EXPECTED_DEV_ROWS = 412
EXPECTED_QA_ROWS = 412
EXPECTED_DEV_CORRUPT_BARCODES = 377  # `01` §3
EXPECTED_DEV_DISTINCT_BARCODES = 98  # `01` §3
EXPECTED_CHAR_VALUE_LIST_ROWS = 195  # `01` §1
EXPECTED_CHAR_GUIDELINES_ROWS = 196  # `01` §1
EXPECTED_CATEGORY = "ORAL HEALTH"  # `01` §7 — single-category problem
EXPECTED_CHARACTERISTIC_COUNT = 13  # `01` §7

# The submission contract (`01` §2). Derived from `OutputRow` rather than
# retyped, so the workbook, the contract and the assembler cannot drift apart
# without something failing loudly here.
EXPECTED_HEADER: list[str] = list(OutputRow.model_fields)

_INPUT_COLUMNS = ("ITEM_CODE", "NAN_KEY", "COUNTRY", "RETAILER_DESC", "RETAILER", "BRAND")


def _worksheet(workbook_path: Path, sheet: str) -> tuple[openpyxl.Workbook, Worksheet]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    if sheet not in workbook.sheetnames:
        raise DatasetDriftError(
            f"sheet {sheet!r} is absent from {workbook_path}; "
            f"present sheets are {workbook.sheetnames!r}. `01` §1 documents the expected set — "
            f"re-verify 01-dataset-contract.md before proceeding."
        )
    return workbook, workbook[sheet]


def read_header(workbook_path: Path, sheet: str) -> list[str]:
    """The sheet's header row, in order. `specs/loader.md` §10."""
    workbook, worksheet = _worksheet(workbook_path, sheet)
    try:
        header = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    finally:
        workbook.close()
    return [str(name) for name in header if name is not None]


def load_qa_header(workbook_path: Path) -> list[str]:
    """Capture and validate the `qa` header — the submission contract.

    A silently reordered submission scores zero regardless of correct values
    (`03` §3), so this asserts name *and* order, not set membership.
    """
    header = read_header(workbook_path, "qa")
    if header != EXPECTED_HEADER:
        raise DatasetDriftError(
            f"qa header does not match the submission contract.\n"
            f"  expected ({len(EXPECTED_HEADER)}): {EXPECTED_HEADER}\n"
            f"  actual   ({len(header)}): {header}\n"
            f"This is `01` §2 / `03` §3's OutputRow field order. A reordered or renamed "
            f"column means either the workbook changed or OutputRow drifted — resolve which "
            f"before proceeding; do not edit the expectation to match."
        )
    return header


def _read_external_codes(workbook_path: Path, sheet: str, expected_rows: int) -> list[str | None]:
    """`EXTERNAL_CODE` via openpyxl, branching on raw cell type.

    `specs/loader.md` §1: this column is the one place raw cell type matters.
    pandas has already inferred a dtype by the time you could look, which
    loses the numeric/text distinction that tells corrupt `dev` values apart
    from intact `qa` ones, and drops the literal leading apostrophe some text
    cells carry.
    """
    workbook, worksheet = _worksheet(workbook_path, sheet)
    try:
        header = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        if "EXTERNAL_CODE" not in header:
            raise DatasetDriftError(f"sheet {sheet!r} has no EXTERNAL_CODE column (`01` §2)")
        index = header.index("EXTERNAL_CODE")
        values: list[str | None] = []
        for row in worksheet.iter_rows(min_row=2, max_row=expected_rows + 1):
            cell = row[index]
            value = cell.value
            if value is None:
                values.append(None)
            elif cell.data_type == "n":
                # Narrowed explicitly rather than trusting data_type alone.
                # openpyxl's 'n' covers anything Excel stores numerically, and
                # int() on a date/time/Decimal would either raise something
                # unhelpful or silently truncate. A barcode cell that is not a
                # real number is a defect worth naming, not coercing (`04` §4).
                if not isinstance(value, int | float):
                    raise DatasetSchemaError(
                        f"{sheet} row {cell.row}: EXTERNAL_CODE has numeric cell type but a "
                        f"{type(value).__name__} value ({value!r}). `01` §3 documents this "
                        f"column as numeric-or-text only."
                    )
                values.append(str(int(value)))
            else:
                values.append(str(value).lstrip("'").strip())
    finally:
        workbook.close()
    return values


def _load_retailer_names(retailers_path: Path) -> dict[str, str]:
    """`config/retailers.yaml` -> {retailer_raw: name}. Internal to this module.

    Kept internal deliberately: `04` §3 forbids a bare dict crossing a module
    boundary, and this table has no reason to.
    """
    raw = yaml.safe_load(retailers_path.read_text(encoding="utf-8"))
    if not isinstance(raw, dict):
        raise DatasetSchemaError(f"{retailers_path} did not parse to a mapping")
    names: dict[str, str] = {}
    for retailer_raw, entry in raw.items():
        if not isinstance(entry, dict) or not entry.get("name"):
            raise DatasetSchemaError(
                f"{retailers_path}: entry {retailer_raw!r} has no `name`. "
                f"`name` is complete-or-nothing for all 50 entries (specs/loader.md §4); "
                f"`domain` may stay partial until P7, `name` may not."
            )
        names[str(retailer_raw)] = str(entry["name"])
    return names


def _assert_input_schema(frame: pd.DataFrame, sheet: str, expected_rows: int) -> None:
    """`05` §5's schema-drift guardrail, generalized past the corruption count.

    `specs/loader.md` §11 asks for column-set and shape assertions alongside
    the corruption count, on the same principle: assert a known fingerprint,
    fail loudly on drift.
    """
    actual = list(frame.columns)
    if actual != EXPECTED_HEADER:
        raise DatasetDriftError(
            f"{sheet} column set/order changed.\n  expected: {EXPECTED_HEADER}\n"
            f"  actual  : {actual}\nRe-verify 01-dataset-contract.md §2."
        )
    if len(frame) != expected_rows:
        raise DatasetDriftError(
            f"{sheet} has {len(frame)} rows, expected {expected_rows} (`01` §1). "
            f"The source file changed — re-verify 01-dataset-contract.md."
        )
    for column in _INPUT_COLUMNS:
        null_count = int(frame[column].isna().sum())
        if null_count:
            raise DatasetDriftError(
                f"{sheet}.{column} has {null_count} null values; `01` §6 records it as "
                f"fully populated. Re-verify 01-dataset-contract.md."
            )


def load_rows(
    workbook_path: Path,
    sheet: str,
    retailers_path: Path,
) -> list[RawRow]:
    """One `RawRow` per data row, in source order. `specs/loader.md` §2–§6.

    Row order matches the sheet exactly — no sorting, no groupby reordering.
    """
    expected_rows = EXPECTED_DEV_ROWS if sheet == "dev" else EXPECTED_QA_ROWS
    frame = pd.read_excel(workbook_path, sheet_name=sheet, dtype=str)
    _assert_input_schema(frame, sheet, expected_rows)

    external_codes = _read_external_codes(workbook_path, sheet, expected_rows)
    if len(external_codes) != len(frame):
        raise DatasetDriftError(
            f"{sheet}: openpyxl read {len(external_codes)} EXTERNAL_CODE cells but pandas read "
            f"{len(frame)} rows. The two readers disagree about the sheet's extent; row "
            f"alignment cannot be trusted."
        )

    retailer_names = _load_retailer_names(retailers_path)
    rows: list[RawRow] = []

    for position, (record, external_code) in enumerate(
        zip(frame.to_dict("records"), external_codes, strict=True)
    ):
        nan_key = int(str(record["NAN_KEY"]))
        barcode, barcode_raw, barcode_corrupt = parse_barcode(external_code)
        if barcode_corrupt:
            log.warning(
                "barcode_corrupt_rounded",
                sheet=sheet,
                nan_key=nan_key,
                barcode_raw=barcode_raw,
            )

        brand_raw, brand_encoding_suspect = repair_encoding(str(record["BRAND"]))
        brand, brand_owner = parse_brand(brand_raw)

        desc_repaired, desc_encoding_suspect = repair_encoding(str(record["RETAILER_DESC"]))
        desc_raw = collapse_whitespace(desc_repaired)

        if brand_encoding_suspect or desc_encoding_suspect:
            log.info(
                "encoding_repaired",
                sheet=sheet,
                nan_key=nan_key,
                brand_before=str(record["BRAND"]) if brand_encoding_suspect else None,
                brand_after=brand_raw if brand_encoding_suspect else None,
                desc_before=str(record["RETAILER_DESC"]) if desc_encoding_suspect else None,
                desc_after=desc_repaired if desc_encoding_suspect else None,
            )

        retailer_raw = str(record["RETAILER"])
        if retailer_raw not in retailer_names:
            raise RetailerNotMappedError(
                f"{sheet} row {position + 2}: RETAILER {retailer_raw!r} has no entry in "
                f"{retailers_path}. Add a hand-reviewed entry — do not fall back to the raw "
                f"string (specs/loader.md §4; an unmapped retailer passing through silently "
                f"is the failure class `01` §12 documents)."
            )

        rows.append(
            RawRow(
                nan_key=nan_key,
                item_code=int(str(record["ITEM_CODE"])),
                barcode=barcode,
                barcode_raw=barcode_raw,
                barcode_corrupt=barcode_corrupt,
                brand_raw=brand_raw,
                brand=brand,
                brand_owner=brand_owner,
                brand_encoding_suspect=brand_encoding_suspect,
                retailer_raw=retailer_raw,
                retailer=retailer_names[retailer_raw],
                countries=parse_countries(str(record["COUNTRY"])),
                desc_raw=desc_raw,
                desc_encoding_suspect=desc_encoding_suspect,
            )
        )

    if sheet == "dev":
        _assert_dev_barcode_fingerprint(rows)
    return rows


def _assert_dev_barcode_fingerprint(rows: list[RawRow]) -> None:
    """The literal P2 gate — `04` §1, `specs/loader.md` §11."""
    corrupt_count = sum(row.barcode_corrupt for row in rows)
    if corrupt_count != EXPECTED_DEV_CORRUPT_BARCODES:
        raise DatasetDriftError(
            f"dev barcode corruption count is {corrupt_count}, expected "
            f"{EXPECTED_DEV_CORRUPT_BARCODES} (01-dataset-contract.md §3). The source file "
            f"changed — re-verify 01-dataset-contract.md before proceeding, don't just update "
            f"this number."
        )
    distinct = len({row.barcode_raw for row in rows if row.barcode_raw is not None})
    if distinct != EXPECTED_DEV_DISTINCT_BARCODES:
        raise DatasetDriftError(
            f"dev has {distinct} distinct EXTERNAL_CODE values, expected "
            f"{EXPECTED_DEV_DISTINCT_BARCODES} (01-dataset-contract.md §3). Re-verify "
            f"01-dataset-contract.md before proceeding."
        )


def _parse_allowed_values(raw: object, module: str, characteristic: str) -> list[str]:
    """`possible_values` is a Python list literal stored as text (`01` §7)."""
    try:
        parsed = ast.literal_eval(str(raw))
    except (ValueError, SyntaxError) as exc:
        raise DatasetSchemaError(
            f"char_value_list possible_values for ({module!r}, {characteristic!r}) is not a "
            f"parseable Python literal: {raw!r}"
        ) from exc
    if not isinstance(parsed, list) or not all(isinstance(value, str) for value in parsed):
        raise DatasetSchemaError(
            f"char_value_list possible_values for ({module!r}, {characteristic!r}) parsed to "
            f"{type(parsed).__name__}, expected list[str]: {raw!r}"
        )
    return [str(value) for value in parsed]


def _parse_binary(raw: object, module: str, characteristic: str) -> bool:
    text = str(raw).strip().upper()
    if text == "Y":
        return True
    if text == "N":
        return False
    raise DatasetSchemaError(
        f"char_value_list binary for ({module!r}, {characteristic!r}) is {raw!r}, expected "
        f"'Y' or 'N' (`01` §7)"
    )


def _parse_open_close(
    raw: object, module: str, characteristic: str
) -> Literal["Close", "Open-ended"]:
    """Returns the Literal type directly so callers need no cast or suppression."""
    text = str(raw).strip()
    if text == "Close":
        return "Close"
    if text == "Open-ended":
        return "Open-ended"
    raise DatasetSchemaError(
        f"char_value_list open_close for ({module!r}, {characteristic!r}) is {raw!r}, "
        f"expected 'Close' or 'Open-ended' (`01` §7)"
    )


def load_characteristic_rules(workbook_path: Path) -> list[CharacteristicRule]:
    """`char_value_list` -> the applicability + vocabulary table. §8."""
    frame = pd.read_excel(workbook_path, sheet_name="char_value_list", dtype=str)
    if len(frame) != EXPECTED_CHAR_VALUE_LIST_ROWS:
        raise DatasetDriftError(
            f"char_value_list has {len(frame)} rows, expected "
            f"{EXPECTED_CHAR_VALUE_LIST_ROWS} (`01` §1). Re-verify 01-dataset-contract.md."
        )

    rules: list[CharacteristicRule] = []
    for record in frame.to_dict("records"):
        category = str(record["category"]).strip()
        if category != EXPECTED_CATEGORY:
            raise DatasetSchemaError(
                f"char_value_list category is {category!r}, expected {EXPECTED_CATEGORY!r}. "
                f"Single-category scope is a standing assumption across `03`/`04`/`05` — a "
                f"second value means the file changed underneath the whole design."
            )
        module = str(record["module"])
        characteristic = normalize_characteristic_name(str(record["characteristic"]))
        rules.append(
            CharacteristicRule(
                module=module,
                characteristic=characteristic,
                open_close=_parse_open_close(record["open_close"], module, characteristic),
                binary=_parse_binary(record["binary"], module, characteristic),
                allowed_values=_parse_allowed_values(
                    record["possible_values"], module, characteristic
                ),
            )
        )
    _assert_characteristic_names_bijective(
        {rule.characteristic for rule in rules}, "char_value_list"
    )
    return rules


def load_characteristic_guidelines(workbook_path: Path) -> list[CharacteristicGuideline]:
    """`char_guidelines` -> per-(module, characteristic) business rules. §9."""
    frame = pd.read_excel(workbook_path, sheet_name="char_guidelines", dtype=str)
    if len(frame) != EXPECTED_CHAR_GUIDELINES_ROWS:
        raise DatasetDriftError(
            f"char_guidelines has {len(frame)} rows, expected "
            f"{EXPECTED_CHAR_GUIDELINES_ROWS} (`01` §1). Re-verify 01-dataset-contract.md."
        )

    guidelines = [
        CharacteristicGuideline(
            module=str(record["MODULE NAME"]),
            characteristic=normalize_characteristic_name(str(record["CHARACTERISTICS NAME"])),
            guideline_text=str(record["Guidelines"]),
        )
        for record in frame.to_dict("records")
    ]
    _assert_characteristic_names_bijective(
        {guideline.characteristic for guideline in guidelines}, "char_guidelines"
    )
    return guidelines


def _assert_characteristic_names_bijective(normalized: set[str], source: str) -> None:
    """§7: total and bijective against the 13 `dev`/`qa` columns, after aliasing."""
    expected = {name for name in EXPECTED_HEADER if name.startswith("GLOBAL_")}
    if len(expected) != EXPECTED_CHARACTERISTIC_COUNT:
        raise DatasetSchemaError(
            f"expected {EXPECTED_CHARACTERISTIC_COUNT} GLOBAL_ columns in the header, "
            f"found {len(expected)}"
        )
    if normalized != expected:
        raise DatasetSchemaError(
            f"{source} characteristic names do not map onto the dev/qa columns bijectively.\n"
            f"  unmatched in {source}: {sorted(normalized - expected)}\n"
            f"  unmatched in dev/qa  : {sorted(expected - normalized)}\n"
            f"See `01` §8 / specs/loader.md §7 — a name needing an alias belongs in "
            f"CHARACTERISTIC_NAME_ALIASES, not in a cleverer regex."
        )


def assert_guideline_modules_subset(
    guidelines: list[CharacteristicGuideline],
    rules: list[CharacteristicRule],
) -> None:
    """§9: `char_guidelines` modules must be a subset of `char_value_list`'s."""
    guideline_modules = {guideline.module for guideline in guidelines}
    rule_modules = {rule.module for rule in rules}
    orphans = guideline_modules - rule_modules
    if orphans:
        raise DatasetSchemaError(
            f"char_guidelines references {len(orphans)} module(s) absent from char_value_list: "
            f"{sorted(orphans)}. This is a cross-sheet consistency check the organizers' own "
            f"files should satisfy — surface it, don't absorb it silently."
        )


def applicable_characteristics(rules: list[CharacteristicRule], module: str) -> list[str]:
    """Which characteristics apply to `module`. `specs/loader.md` §8.

    A plain function, not a contract type. Everything not returned here must
    be left empty for that module — guessing a non-applicable value is a wrong
    answer, not a partial-credit one (`00`, `01` §7).
    """
    return [rule.characteristic for rule in rules if rule.module == module]


def characteristic_rule(
    rules: list[CharacteristicRule], module: str, characteristic: str
) -> CharacteristicRule | None:
    """The single rule for a (module, characteristic) pair, or None."""
    for rule in rules:
        if rule.module == module and rule.characteristic == characteristic:
            return rule
    return None
